import os
import docx
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


def generate_sample_pdf(filepath: str):
    """Generates a comprehensive senior software engineer resume as a PDF."""
    doc = SimpleDocTemplate(filepath, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1A365D'),
        spaceAfter=4
    )
    
    subtitle_style = ParagraphStyle(
        'SubtitleStyle',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#4A5568'),
        spaceAfter=12
    )
    
    heading_style = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#2B6CB0'),
        spaceBefore=10,
        spaceAfter=6,
        keepWithNext=True
    )
    
    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor('#2D3748'),
        spaceAfter=4
    )
    
    bullet_style = ParagraphStyle(
        'Bullet',
        parent=styles['Normal'],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#2D3748'),
        leftIndent=15,
        spaceAfter=2
    )

    story = []

    # Header
    story.append(Paragraph("Alex Mercer", title_style))
    story.append(Paragraph(
        "Email: alex.mercer@example.com | Phone: +1 (555) 234-5678 | Location: San Francisco, CA<br/>"
        "LinkedIn: https://linkedin.com/in/alex-mercer | GitHub: https://github.com/alexmercer-dev",
        subtitle_style
    ))
    story.append(Spacer(1, 6))

    # Professional Summary
    story.append(Paragraph("PROFESSIONAL SUMMARY", heading_style))
    story.append(Paragraph(
        "Passionate Senior Full-Stack Software Engineer with over 6 years of experience architecting and building "
        "scalable cloud-native web applications, distributed systems, and modern microservices. Proven track record "
        "in optimizing performance, leading cross-functional teams, and deploying robust machine learning pipelines.",
        body_style
    ))
    story.append(Spacer(1, 6))

    # Skills
    story.append(Paragraph("SKILLS", heading_style))
    story.append(Paragraph(
        "<b>Technical Skills:</b> Python, FastAPI, Django, TypeScript, React, Next.js, Node.js, PostgreSQL, Redis, Docker, Kubernetes, AWS (ECS, S3, RDS), GraphQL, Git, CI/CD",
        body_style
    ))
    story.append(Paragraph(
        "<b>Soft Skills:</b> Agile Leadership, System Design Mentorship, Technical Writing, Problem Solving, Cross-Functional Communication",
        body_style
    ))
    story.append(Spacer(1, 6))

    # Work Experience
    story.append(Paragraph("WORK EXPERIENCE", heading_style))
    
    story.append(Paragraph("<b>Staff Software Engineer</b> - CloudScale Technologies (San Francisco, CA)", body_style))
    story.append(Paragraph("<i>Jan 2022 - Present</i>", subtitle_style))
    story.append(Paragraph("• Architected high-throughput microservices using FastAPI and Redis, decreasing API response latency by 45% for 2M+ daily active users.", bullet_style))
    story.append(Paragraph("• Led migration of legacy monolith into containerized Kubernetes services on AWS EKS, improving deployment cycle time from bi-weekly to daily.", bullet_style))
    story.append(Paragraph("• Mentored an engineering squad of 8 frontend and backend engineers, conducting code reviews and system design workshops.", bullet_style))
    story.append(Spacer(1, 4))

    story.append(Paragraph("<b>Senior Backend Engineer</b> - DataPulse Analytics (New York, NY)", body_style))
    story.append(Paragraph("<i>Jun 2019 - Dec 2021</i>", subtitle_style))
    story.append(Paragraph("• Designed real-time event streaming pipeline processing 10,000 events/sec using Apache Kafka and Python.", bullet_style))
    story.append(Paragraph("• Re-engineered database indexing and query patterns on PostgreSQL, reducing slow query counts by 60%.", bullet_style))
    story.append(Paragraph("• Implemented OAuth2 and JWT authentication mechanisms adhering to SOC2 compliance standards.", bullet_style))
    story.append(Spacer(1, 4))

    story.append(Paragraph("<b>Software Developer</b> - NexaSoft Innovations (Boston, MA)", body_style))
    story.append(Paragraph("<i>Jul 2017 - May 2019</i>", subtitle_style))
    story.append(Paragraph("• Built dynamic single-page web applications with React and Redux, integrated with RESTful backends.", bullet_style))
    story.append(Paragraph("• Developed unit and end-to-end automated test suites with Pytest and Cypress, raising code coverage from 55% to 88%.", bullet_style))
    story.append(Spacer(1, 6))

    # Education
    story.append(Paragraph("EDUCATION", heading_style))
    story.append(Paragraph("<b>B.S. in Computer Science</b>", body_style))
    story.append(Paragraph("University of California, Berkeley (2013 - 2017) | Grade: GPA 3.85 / 4.0", subtitle_style))
    story.append(Spacer(1, 6))

    # Projects
    story.append(Paragraph("KEY PROJECTS", heading_style))
    story.append(Paragraph("<b>FlowTrack - Workflow Automation Engine</b>", body_style))
    story.append(Paragraph("An open-source distributed task runner built with Python, Celery, and WebSockets enabling real-time DAG execution visualization.", bullet_style))
    story.append(Paragraph("<i>Tech Used: Python, Celery, Redis, React, WebSockets</i>", subtitle_style))
    story.append(Spacer(1, 4))

    story.append(Paragraph("<b>DocuQuery - Semantic Document Assistant</b>", body_style))
    story.append(Paragraph("RAG-based search engine for enterprise documentation leveraging Claude API and pgvector for semantic retrieval.", bullet_style))
    story.append(Paragraph("<i>Tech Used: Anthropic Claude, Python, LangChain, PostgreSQL pgvector</i>", subtitle_style))
    story.append(Spacer(1, 6))

    # Certifications
    story.append(Paragraph("CERTIFICATIONS", heading_style))
    story.append(Paragraph("• AWS Certified Solutions Architect - Associate (Amazon Web Services, 2022)", body_style))
    story.append(Paragraph("• Certified Kubernetes Administrator - CKA (Cloud Native Computing Foundation, 2023)", body_style))

    doc.build(story)
    print(f"Generated PDF: {filepath}")


def generate_sample_docx(filepath: str):
    """Generates an entry-level / partial resume as DOCX to test missing fields & confidence."""
    doc = docx.Document()
    
    # Title / Name
    doc.add_heading("Taylor Brooks", level=0)
    
    # Missing phone number and certifications intentionally
    p_contact = doc.add_paragraph()
    p_contact.add_run("Email: taylor.brooks@example.com | Location: Seattle, WA\n")
    p_contact.add_run("GitHub: https://github.com/taylorbrooks-ds")

    doc.add_heading("Objective", level=1)
    doc.add_paragraph(
        "Motivated Junior Data Scientist with strong foundational skills in Python, machine learning, and statistical analysis. "
        "Eager to contribute to high-impact predictive modeling and data visualization projects."
    )

    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Technical: Python, Pandas, Scikit-Learn, SQL, Matplotlib, Seaborn, Git")
    doc.add_paragraph("Soft: Teamwork, Critical Thinking, Analytical Mindset")

    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Data Science Intern - Insights Corp (Seattle, WA)")
    p_dates = doc.add_paragraph("Jun 2023 - Sep 2023")
    doc.add_paragraph("• Built customer churn classification model achieving 82% ROC-AUC using Random Forests.")
    doc.add_paragraph("• Cleaned and preprocessed customer telemetry dataset containing 500k+ rows with Pandas.")

    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.S. in Statistics - University of Washington (2020 - 2024)")

    doc.add_heading("Academic Projects", level=1)
    doc.add_paragraph("Housing Price Predictor: Regression modeling on King County housing data using XGBoost and Streamlit.")

    doc.save(filepath)
    print(f"Generated DOCX: {filepath}")


def generate_dummy_scanned_pdf(filepath: str):
    """Generates an empty / text-less PDF to test scanned document error handling."""
    doc = SimpleDocTemplate(filepath, pagesize=letter)
    # Build with an empty spacer only, no text paragraphs
    doc.build([Spacer(1, 100)])
    print(f"Generated Scanned Dummy PDF: {filepath}")


if __name__ == "__main__":
    os.makedirs("sample_resumes", exist_ok=True)
    generate_sample_pdf("sample_resumes/alex_mercer_senior_engineer.pdf")
    generate_sample_docx("sample_resumes/taylor_brooks_junior_ds.docx")
    generate_dummy_scanned_pdf("sample_resumes/empty_scanned_sample.pdf")
    print("All sample resumes created successfully in sample_resumes/ folder!")
