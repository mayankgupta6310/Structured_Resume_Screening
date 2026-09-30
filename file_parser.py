import io
import os
from typing import Union
import docx
import pdfplumber


class FileParserError(Exception):
    """Base exception for file parsing errors."""
    pass


class ScannedPDFError(FileParserError):
    """Raised when a PDF has no extractable text, indicating it may be a scanned image."""
    pass


class EmptyDocumentError(FileParserError):
    """Raised when an uploaded document contains no readable text."""
    pass


def extract_text_from_pdf(file_source: Union[str, bytes, io.BytesIO]) -> str:
    """
    Extracts text from a PDF file using pdfplumber.
    Supports file path (str), raw bytes, or BytesIO buffer.
    
    Raises:
        ScannedPDFError: If no text could be extracted (likely scanned).
        EmptyDocumentError: If document is empty.
        FileParserError: For other parsing errors.
    """
    try:
        if isinstance(file_source, str):
            pdf_ctx = pdfplumber.open(file_source)
        elif isinstance(file_source, bytes):
            pdf_ctx = pdfplumber.open(io.BytesIO(file_source))
        elif isinstance(file_source, io.BytesIO):
            file_source.seek(0)
            pdf_ctx = pdfplumber.open(file_source)
        else:
            # e.g., Streamlit UploadedFile
            pdf_ctx = pdfplumber.open(file_source)

        text_parts = []
        extracted_links = []
        with pdf_ctx as pdf:
            if len(pdf.pages) == 0:
                raise EmptyDocumentError("The uploaded PDF has 0 pages.")

            for page_idx, page in enumerate(pdf.pages, start=1):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text.strip())

                # Also extract text inside tables if any
                tables = page.extract_tables()
                for table in tables:
                    for row in table:
                        row_text = " | ".join([cell.strip() for cell in row if cell and cell.strip()])
                        if row_text and row_text not in text_parts:
                            text_parts.append(row_text)

                # Extract hyperlink annotations from PDF
                hyperlinks = getattr(page, "hyperlinks", []) or []
                for link in hyperlinks:
                    uri = link.get("uri")
                    if not uri:
                        continue
                    if isinstance(uri, bytes):
                        uri = uri.decode("utf-8", errors="ignore")
                    uri = str(uri).strip()
                    if not uri:
                        continue

                    text = link.get("contents")
                    if isinstance(text, bytes):
                        text = text.decode("utf-8", errors="ignore")
                    if not text:
                        try:
                            x0 = link.get("x0", 0)
                            top = link.get("top", link.get("doctop", 0))
                            x1 = link.get("x1", 0)
                            bottom = link.get("bottom", 0)
                            if x1 > x0 and bottom > top:
                                cropped = page.crop((x0, top, x1, bottom))
                                cropped_text = cropped.extract_text()
                                if cropped_text:
                                    text = cropped_text.strip()
                        except Exception:
                            pass

                    visible_label = str(text).strip() if text else "Link"
                    link_line = f"Link: [{visible_label}] -> [{uri}]"
                    if link_line not in extracted_links:
                        extracted_links.append(link_line)

        full_text = "\n\n".join(text_parts).strip()

        if not full_text:
            raise ScannedPDFError(
                "No readable text found in the PDF. This document may be a scanned image or empty. "
                "Please provide a text-based PDF or DOCX file."
            )

        # Basic check: if full_text is too short (e.g. less than 20 characters)
        if len(full_text) < 20:
            raise EmptyDocumentError(
                "The PDF contains almost no text (less than 20 characters). Please provide a complete resume."
            )

        if extracted_links:
            full_text += "\n\nExtracted Hyperlinks:\n" + "\n".join(extracted_links)

        return full_text

    except (ScannedPDFError, EmptyDocumentError):
        raise
    except Exception as e:
        raise FileParserError(f"Failed to parse PDF: {str(e)}") from e


def extract_text_from_docx(file_source: Union[str, bytes, io.BytesIO]) -> str:
    """
    Extracts text from a DOCX file using python-docx.
    Supports file path (str), raw bytes, or BytesIO buffer.
    
    Raises:
        EmptyDocumentError: If document contains no readable text.
        FileParserError: For other parsing errors.
    """
    try:
        if isinstance(file_source, str):
            doc = docx.Document(file_source)
        elif isinstance(file_source, bytes):
            doc = docx.Document(io.BytesIO(file_source))
        elif isinstance(file_source, io.BytesIO):
            file_source.seek(0)
            doc = docx.Document(file_source)
        else:
            doc = docx.Document(file_source)

        text_parts = []

        # Extract paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                text_parts.append(text)

        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                # remove adjacent duplicates caused by merged cells
                cleaned_cells = []
                for c in row_cells:
                    if not cleaned_cells or c != cleaned_cells[-1]:
                        cleaned_cells.append(c)
                if cleaned_cells:
                    text_parts.append(" | ".join(cleaned_cells))

        # Extract hyperlinks from DOCX
        docx_links = []
        try:
            for p in doc.paragraphs:
                for child in p._p:
                    if child.tag.endswith('hyperlink'):
                        r_id = child.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                        if r_id and r_id in doc.part.rels:
                            url = doc.part.rels[r_id].target_ref
                            label = "".join(node.text for node in child.iter() if node.text).strip()
                            if url:
                                link_line = f"Link: [{label or 'Link'}] -> [{url}]"
                                if link_line not in docx_links:
                                    docx_links.append(link_line)
        except Exception:
            pass

        full_text = "\n\n".join(text_parts).strip()

        if not full_text:
            raise EmptyDocumentError("No readable text found in the DOCX file.")

        if len(full_text) < 20:
            raise EmptyDocumentError(
                "The DOCX file contains almost no text (less than 20 characters). Please provide a complete resume."
            )

        if docx_links:
            full_text += "\n\nExtracted Hyperlinks:\n" + "\n".join(docx_links)

        return full_text

    except EmptyDocumentError:
        raise
    except Exception as e:
        raise FileParserError(f"Failed to parse DOCX: {str(e)}") from e


def parse_resume_file(file_source: Union[str, bytes, io.BytesIO], filename: str) -> str:
    """
    Unified entry point for parsing resume files (PDF or DOCX).
    
    Args:
        file_source: File path, bytes, or BytesIO stream.
        filename: Name of the file (used to detect extension).
        
    Returns:
        Clean plain text of the resume.
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_source)
    elif ext in [".docx", ".doc"]:
        if ext == ".doc":
            raise FileParserError(
                "Legacy .doc format is not supported. Please save/convert the document to .docx or .pdf."
            )
        return extract_text_from_docx(file_source)
    else:
        raise FileParserError(f"Unsupported file format '{ext}'. Please upload a .pdf or .docx file.")
